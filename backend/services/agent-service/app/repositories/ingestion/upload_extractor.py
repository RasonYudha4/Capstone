# app/repositories/ingestion/upload_extractor.py
from __future__ import annotations
from pathlib import Path

from app.repositories.ingestion.parser import (
    _extract_pymupdf4llm, _extract_pdfplumber, _extract_pypdf, _strip_watermarks,
)
from docx import Document
from app.core.logger import get_logger

log = get_logger("upload_extractor")
MAX_CHARS = 8000

def extract_text_from_upload(file_path: str | Path, content_type: str) -> str | None:
    path = Path(file_path)

    if path.suffix.lower() == ".pdf":
        text = _extract_pdf(path)
    elif path.suffix.lower() == ".docx":
        text = _extract_docx(path)
    else:
        log.warning("unsupported upload type: %s", content_type)
        return None

    if text is None or not text.strip():
        return None  # same "give up" behavior as parser.py — no OCR fallback

    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n...[dipotong karena terlalu panjang]"
    return text.strip()


def _extract_pdf(path: Path) -> str | None:
    """Reuses the exact same fallback chain as parser.py's _parse_pdf."""
    for extractor in (_extract_pymupdf4llm, _extract_pdfplumber, _extract_pypdf):
        try:
            _, full_text = extractor(path)
            if full_text.strip():
                return full_text
        except Exception as exc:
            log.warning("%s failed for upload %s: %s", extractor.__name__, path.name, exc)
    return None  # scanned PDF → fails, exactly like ingestion does


def _extract_docx(path: Path) -> str:
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)